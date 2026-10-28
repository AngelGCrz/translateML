"""Lectura de Multidioma.xlsx local y generación de ENG/ESP/DEU de principio a fin.

Requiere openpyxl (requirements.txt). Sin él, estas pruebas se omiten con un aviso.
"""
import csv
import io
import json
import os
import tempfile
import unittest
from contextlib import redirect_stderr, redirect_stdout
from pathlib import Path

from translateml.cli import main
from translateml.sheets.errors import SheetsReadError, SheetsStructureError
from tests.helpers import REAL_TEMPLATES, test_settings, write_synthetic
from tests.xlsx_builder import E, F, google_formula, write_xlsx

try:
    import openpyxl  # noqa: F401
    from translateml.sheets.xlsx import XlsxWorkbook
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

GAME = "g999"
ES = {
    "T0001": "REGLAS DEL JUEGO", "T0002": "Juega hasta 50 líneas.",
    "T0003": "##sym_11.png##paga en cualquier posición.", "T0004": "Texto del jackpot",
    "T0005": "REGLAS DEL JUEGO", "T0006": "Línea anidada", "T0007": "ÁRBOL DE\nLUZ",
    "T0008": "DE UN TIPO", "T0009": "DEL MISMO TIPO",
}


KEY_MODE = Path(__file__).resolve().parent / "fixtures" / "sheets-key-mode.json"


def run(*args):
    if args and args[0] == "xlsx" and "--sheets-config" not in args:
        args = ("xlsx", "--sheets-config", str(KEY_MODE)) + tuple(args[1:])
    with test_settings(), redirect_stdout(io.StringIO()) as out, redirect_stderr(io.StringIO()) as err:
        code = main(list(args))
    return code, out.getvalue() + err.getvalue()


def read_csv(file):
    with Path(file).open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.reader(handle))


@unittest.skipUnless(HAS_OPENPYXL, "openpyxl no está instalado")
class XlsxReaderTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.dir = Path(self.tmp.name)

    def test_reads_saved_values_formulas_and_blank_cells(self):
        file = write_xlsx(self.dir / "Multidioma.xlsx", {
            "ENG": [["KEY", "TEXT"], ["g:T0001", "  two\nlines "]],
            "ESP": [["TEXT"], [F(google_formula('GOOGLETRANSLATE(ENG!B2,"en","es")', "dos"), "dos")],
                    [F("=ENG!B3", None)], [E("#N/A")], [25], [2.5], ["   "], [""]],
        })
        before = file.read_bytes()
        workbook = XlsxWorkbook(file)
        self.assertEqual(workbook.tabs(), ["ENG", "ESP"])
        self.assertEqual(workbook.read("ENG")[1], ["g:T0001", "  two\nlines "])
        self.assertEqual([r[0] if r else "" for r in workbook.read("ESP")], ["TEXT", "dos", "", "#N/A", "25", "2.5", "   ", ""])
        formulas = workbook.read("ESP", formulas=True)
        self.assertTrue(formulas[1][0].startswith("=IFERROR(__xludf.DUMMYFUNCTION("))
        self.assertEqual(formulas[2][0], "=ENG!B3")
        self.assertEqual(workbook.missing_values("ESP"), {(3, 0)})
        self.assertEqual(file.read_bytes(), before, "el .xlsx no se modifica")

    def test_missing_tab_and_invalid_files(self):
        file = write_xlsx(self.dir / "Multidioma.xlsx", {"ENG": [["TEXT"]]})
        with self.assertRaisesRegex(SheetsStructureError, "DEU"):
            XlsxWorkbook(file).read("DEU")
        with self.assertRaisesRegex(SheetsReadError, "No existe"):
            XlsxWorkbook(self.dir / "nope.xlsx")
        (self.dir / "fake.xlsx").write_text("not a zip", encoding="utf-8")
        with self.assertRaises(SheetsReadError):
            XlsxWorkbook(self.dir / "fake.xlsx")
        (self.dir / "data.csv").write_text("KEY,TEXT", encoding="utf-8")
        with self.assertRaisesRegex(SheetsReadError, "xlsx"):
            XlsxWorkbook(self.dir / "data.csv")


@unittest.skipUnless(HAS_OPENPYXL, "openpyxl no está instalado")
class XlsxFlowTest(unittest.TestCase):
    """extract -> export (claves para ENG) -> Data trabaja en Sheets -> descarga .xlsx -> build."""

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        tmp = Path(self.tmp.name)
        (tmp / "tpl").mkdir()
        self.template = write_synthetic(tmp / "tpl")
        self.out = tmp / "out"
        self.manifest = str(self.out / "manifest.json")
        self.xlsx = tmp / "Multidioma.xlsx"
        self.assertEqual(run("extract", str(self.template), "--out", str(self.out), "--name", GAME)[0], 0)
        self.assertEqual(run("xlsx", "export", self.manifest)[0], 0)
        self.eng = read_csv(self.out / "eng-sheet.csv")

    def workbook(self, esp=ES, deu=None, extra_tabs=None, eng=None):
        """Libro como lo descargaría Data: ENG pegado desde eng-sheet.csv y fórmulas en ESP/DEU."""
        eng = eng or self.eng
        tabs = {"ENG": eng}
        for tab, translations in (("ESP", esp), ("DEU", deu or {})):
            rows = [["TEXT"]]
            for r, (key, _) in enumerate(eng[1:], start=2):
                value = translations.get(key.split(":")[1], "")
                rows.append([F(google_formula(f'GOOGLETRANSLATE(ENG!B{r},"en","{tab.lower()[:2]}")', value), value)])
            tabs[tab] = rows
        tabs.update(extra_tabs or {})
        return write_xlsx(self.xlsx, tabs)

    def build(self, *extra):
        return run("xlsx", "build", self.manifest, "--xlsx", str(self.xlsx), *extra)

    def report(self):
        return json.loads((self.out / "xlsx-build-report.json").read_text(encoding="utf-8"))

    def test_export_has_keys_in_document_order(self):
        self.assertEqual(self.eng[0], ["KEY", "TEXT"])
        self.assertEqual(self.eng[1], [f"{GAME}:T0001", "GAME RULES"])
        self.assertEqual(self.eng[7], [f"{GAME}:T0007", "TREE OF\nLIGHT"])
        self.assertEqual(len(self.eng), 10)

    def test_generates_eng_esp_deu(self):
        self.workbook(deu={"T0001": "SPIELREGELN"})
        code, output = self.build()
        self.assertEqual(code, 0, output)
        self.assertEqual((self.out / "eng" / "info.json").read_bytes(), self.template.read_bytes())
        esp = json.loads((self.out / "esp" / "info.json").read_text(encoding="utf-8"))
        self.assertEqual(esp["rules"]["pages"][0]["items"]["lines"],
                         ["• Juega hasta 50 líneas. ", "", "   ##sym_11.png##paga en cualquier posición.",
                          "{fixedJackpot}• Texto del jackpot", "    6 ##sym_12.png##."])
        self.assertEqual([t["text"] for t in esp["payTables"][0]["titles"]], ["ÁRBOL DE\nLUZ", "DE UN TIPO", "DEL MISMO TIPO"])
        self.assertEqual(esp["payTables"][0]["titles"][0]["tag"], "SCATTER")
        self.assertEqual(esp["fontName"], "altehaasgrotesk bold")
        deu = json.loads((self.out / "deu" / "info.json").read_text(encoding="utf-8"))
        self.assertEqual(deu["rules"]["pages"][0]["titles"][0]["text"], "SPIELREGELN")
        self.assertEqual(deu["payTables"][0]["titles"][1]["text"], "OF A KIND")

        report = self.report()
        self.assertEqual(report["status"], "warnings")
        self.assertEqual(report["extracted_texts"], 9)
        self.assertEqual(report["pending_review"], 3)
        self.assertEqual(report["languages"]["esp"]["applied"], 9)
        self.assertEqual(report["languages"]["deu"]["applied"], 1)
        self.assertEqual(len(report["languages"]["deu"]["missing"]), 8)
        self.assertEqual(read_csv(self.out / "sheets" / "esp.csv")[1], ["T0001", "REGLAS DEL JUEGO", ""])

    def test_lost_marker_keeps_english_and_is_reported(self):
        self.workbook(esp=dict(ES, T0003="## sym_11.png ## paga en cualquier posición."))
        code, output = self.build("--lang", "esp")
        self.assertEqual(code, 0, output)
        esp = json.loads((self.out / "esp" / "info.json").read_text(encoding="utf-8"))
        self.assertEqual(esp["rules"]["pages"][0]["items"]["lines"][2], "   ##sym_11.png##pays anywhere.")
        issues = self.report()["languages"]["esp"]["marker_issues"]
        self.assertEqual([(i["key"], i["kind"], i["blocking"]) for i in issues], [("T0003", "token_mismatch", True)])
        self.assertIn("T0003", self.report()["languages"]["esp"]["missing"])

    def test_missing_tab_fails_that_language_only(self):
        esp = [["TEXT"]] + [[F(f"=ENG!B{r}", ES[k.split(":")[1]])] for r, (k, _) in enumerate(self.eng[1:], start=2)]
        write_xlsx(self.xlsx, {"ENG": self.eng, "ESP": esp})
        code, output = self.build()
        self.assertEqual(code, 1)
        self.assertIn("no tiene la pestaña 'DEU'", output)
        self.assertTrue((self.out / "esp" / "info.json").is_file())
        self.assertFalse((self.out / "deu" / "info.json").exists())
        self.assertEqual(self.report()["status"], "error")

    def test_ambiguous_association_blocks_the_language(self):
        rows = [["TEXT"]] + [[F(f"=ENG!B{r}", ES[k.split(":")[1]])] for r, (k, _) in enumerate(self.eng[1:], start=2)]
        rows[3] = [F("=ENG!B2", "duplicada")]
        write_xlsx(self.xlsx, {"ENG": self.eng, "ESP": rows, "DEU": [["TEXT"]]})
        code, output = self.build("--lang", "esp")
        self.assertEqual(code, 1)
        self.assertIn("ambiguous", output)
        self.assertFalse((self.out / "esp" / "info.json").exists())

    def test_unsaved_formula_values_are_explained(self):
        rows = [["TEXT"]] + [[F(f"=ENG!B{r}", None)] for r in range(2, len(self.eng) + 1)]
        write_xlsx(self.xlsx, {"ENG": self.eng, "ESP": rows, "DEU": [["TEXT"]]})
        code, output = self.build("--lang", "esp")
        self.assertEqual(code, 1)
        self.assertIn("no_cached_value", output)
        self.assertIn("vuelve a descargar el archivo", output)
        code, output = run("xlsx", "inspect", "--xlsx", str(self.xlsx))
        self.assertEqual(code, 0)
        self.assertIn("sin valor guardado", output)

    def test_eng_without_key_column_is_a_clear_error(self):
        write_xlsx(self.xlsx, {"ENG": [["TEXT"]] + [[t] for _, t in self.eng[1:]], "ESP": [["TEXT"]], "DEU": [["TEXT"]]})
        code, output = self.build()
        self.assertEqual(code, 2)
        self.assertIn("no tiene la columna 'KEY'", output)

    def test_missing_text_header_is_a_clear_error(self):
        write_xlsx(self.xlsx, {"ENG": [["KEY", "TEXTO"]] + self.eng[1:], "ESP": [["TEXT"]], "DEU": [["TEXT"]]})
        code, output = self.build()
        self.assertEqual(code, 2)
        self.assertIn("no tiene la columna 'TEXT'", output)
        self.assertIn("'TEXTO'", output)

    def test_duplicate_headers_and_keys(self):
        write_xlsx(self.xlsx, {"ENG": [["KEY", "TEXT", "text"]] + self.eng[1:], "ESP": [["TEXT"]], "DEU": [["TEXT"]]})
        code, output = self.build()
        self.assertEqual(code, 2)
        self.assertIn("repetida", output)
        self.workbook(eng=self.eng + [[f"{GAME}:T0002", "Play up to 50 lines."]])
        code, output = self.build("--lang", "esp")
        self.assertEqual(code, 1)
        self.assertIn("duplicate_key", output)

    def test_export_compares_with_downloaded_file(self):
        partial = self.eng[:5] + [[f"{GAME}:T0099", "Removed rule"]]
        partial[2] = [partial[2][0], "Play up to 40 lines."]
        self.workbook(eng=partial)
        code, output = run("xlsx", "export", self.manifest, "--xlsx", str(self.xlsx))
        self.assertEqual(code, 0, output)
        new_rows = read_csv(self.out / "eng-new-rows.csv")
        self.assertEqual([r[0] for r in new_rows[1:]], [f"{GAME}:T000{i}" for i in range(5, 10)])
        report = json.loads((self.out / "eng-sheet-report.json").read_text(encoding="utf-8"))
        self.assertEqual([c["key"] for c in report["changed"]], [f"{GAME}:T0002"])
        self.assertEqual([o["key"] for o in report["obsolete"]], [f"{GAME}:T0099"])

        code, output = self.build("--lang", "esp")
        self.assertEqual(code, 0, output)
        problems = {(p["kind"], p["key"]) for p in self.report()["languages"]["esp"]["problems"]}
        self.assertIn(("eng_out_of_sync", "T0002"), problems)
        self.assertIn("T0006", self.report()["languages"]["esp"]["missing"])

    def test_real_template_round_trip_through_xlsx(self):
        out = Path(self.tmp.name) / "g402"
        self.assertEqual(run("extract", str(REAL_TEMPLATES[0]), "--out", str(out))[0], 0)
        self.assertEqual(run("xlsx", "export", str(out / "manifest.json"))[0], 0)
        eng = read_csv(out / "eng-sheet.csv")
        esp = [["TEXT"]] + [[F(f"=ENG!B{r}", text.upper())] for r, (_, text) in enumerate(eng[1:], start=2)]
        write_xlsx(self.xlsx, {"ENG": eng, "ESP": esp, "DEU": [["TEXT"]]})
        code, output = run("xlsx", "build", str(out / "manifest.json"), "--xlsx", str(self.xlsx), "--lang", "eng,esp")
        self.assertEqual(code, 0, output)
        self.assertEqual((out / "eng" / "info.json").read_bytes(), REAL_TEMPLATES[0].read_bytes())
        original = json.loads(REAL_TEMPLATES[0].read_text(encoding="utf-8"))
        generated = json.loads((out / "esp" / "info.json").read_text(encoding="utf-8"))
        self.assertEqual(generated["payTables"][0]["titles"][0]["text"], "FAIRY")
        self.assertEqual(generated["rules"]["pages"][0]["items"]["lines"][0], "• PLAY UP TO 50 LINES.")
        self.assertEqual(generated["rules"]["settings"], original["rules"]["settings"])
        self.assertEqual(len(generated["rules"]["pages"]), len(original["rules"]["pages"]))



TEXT_MODE = Path(__file__).resolve().parent / "fixtures" / "sheets-text-mode.json"


class TextModeCase(unittest.TestCase):
    """Preparación común: hoja como el Multidioma real: sin cabeceras ni claves, numeración en B, texto en C, rótulos de sección,
    frases completas y fórmulas GOOGLETRANSLATE(ENG!C<fila>) exportadas por Google."""

    ENG = ["PAYTABLE", "TREE OF\nLIGHT", "OF A KIND", None, "GAME RULES", "Play up to 50 lines.",
           "Play up to 50 lines. ##sym_11.png##pays anywhere.", "Jackpot text", "Nested line!"]
    ES = {"TREE OF\nLIGHT": "ÁRBOL DE\nLUZ", "OF A KIND": "DE UN TIPO", "GAME RULES": "REGLAS DEL JUEGO",
          "Play up to 50 lines.": "Juega hasta 50 líneas.", "Jackpot text": "Texto del jackpot",
          "Play up to 50 lines. ##sym_11.png##pays anywhere.": "Frase completa ##sym_11.png##.",
          "Nested line!": "¡Línea!"}

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        tmp = Path(self.tmp.name)
        (tmp / "tpl").mkdir()
        self.template = write_synthetic(tmp / "tpl")
        self.out = tmp / "out"
        self.manifest = str(self.out / "manifest.json")
        self.xlsx = tmp / "Multidioma.xlsx"
        self.assertEqual(run("extract", str(self.template), "--out", str(self.out), "--name", GAME)[0], 0)

    def write(self, es=None):
        es = self.ES if es is None else es
        eng, esp, deu = [], [], []
        number = 0
        for r, text in enumerate(self.ENG, start=1):
            if text is None:
                eng.append([]), esp.append([]), deu.append([])
                continue
            label = text.isupper() and r in (1, 5)
            number = 0 if label else number + 1
            n = "" if label else str(number)
            eng.append(["", n, text])
            if r == 1:
                esp.append(["", n, text]), deu.append(["", n, text])  # rótulo copiado, sin fórmula
                continue
            esp.append(["", n, F(google_formula(f'GOOGLETRANSLATE(ENG!C{r},"en","es")', es.get(text, "")), es.get(text, ""))])
            deu.append(["", n, F(google_formula(f'GOOGLETRANSLATE(ENG!C{r},"en","de")', "DE"), "DE")])
        write_xlsx(self.xlsx, {"ENG": eng, "ESP": esp, "DEU": deu})

    def build(self, *extra):
        return run("xlsx", "--sheets-config", str(TEXT_MODE), "build", self.manifest, "--xlsx", str(self.xlsx), *extra)



@unittest.skipUnless(HAS_OPENPYXL, "openpyxl no está instalado")
class XlsxTextModeFlowTest(TextModeCase):
    def test_generates_by_exact_english_text(self):
        self.write()
        code, output = self.build()
        self.assertEqual(code, 0, output)
        self.assertEqual((self.out / "eng" / "info.json").read_bytes(), self.template.read_bytes())
        esp = json.loads((self.out / "esp" / "info.json").read_text(encoding="utf-8"))
        pages = esp["rules"]["pages"]
        self.assertEqual([pages[0]["titles"][0]["text"], pages[1]["titles"][0]["text"]], ["REGLAS DEL JUEGO"] * 2)
        self.assertEqual(pages[0]["items"]["lines"][0], "• Juega hasta 50 líneas. ")
        self.assertEqual(pages[0]["items"]["lines"][2], "   ##sym_11.png##pays anywhere.")  # frase en 2 líneas: inglés
        self.assertEqual(pages[0]["items"]["lines"][3], "{fixedJackpot}• Texto del jackpot")
        self.assertEqual(pages[1]["blocks"][0]["items"]["lines"][0], "• Nested line")  # inglés distinto: no se usa
        self.assertEqual([t["text"] for t in esp["payTables"][0]["titles"]], ["ÁRBOL DE\nLUZ", "DE UN TIPO", "DE UN TIPO"])

        report = json.loads((self.out / "xlsx-build-report.json").read_text(encoding="utf-8"))
        self.assertEqual(report["status"], "warnings")
        self.assertEqual(report["association"], "text")
        kinds = sorted(p["kind"] for p in report["eng_problems"])
        # La frase "Play up to 50 lines. ##sym_11…" no se une: "Play up to 50 lines." ya tiene fila propia.
        self.assertEqual(kinds, ["not_in_template", "not_in_template", "sentence_overlaps_rows"])
        self.assertEqual(report["languages"]["esp"]["applied"], 7)
        self.assertEqual(sorted(report["languages"]["esp"]["missing"]), ["T0003", "T0006"])
        self.assertEqual([p["kind"] for p in report["languages"]["esp"]["problems"]], ["unlinked"])

    def lines(self, code="esp"):
        return json.loads((self.out / code / "info.json").read_text(encoding="utf-8"))["rules"]["pages"][0]["items"]["lines"]

    def test_sentence_over_several_lines_is_split_or_merged(self):
        sentence = "Play up to 50 lines. ##sym_11.png##pays anywhere."
        self.ENG = ["GAME RULES", sentence, "Jackpot text"]
        self.ES = {"GAME RULES": "REGLAS", "Jackpot text": "Texto del jackpot",
                   sentence: "Juega hasta 50 líneas.\n##sym_11.png##paga en cualquier posición."}
        self.write()
        code, output = self.build("--lang", "esp")
        self.assertEqual(code, 0, output)
        # Mismos saltos de línea que la plantilla: cada trozo en su línea.
        self.assertEqual(self.lines()[:4], ["• Juega hasta 50 líneas. ", "", "   ##sym_11.png##paga en cualquier posición.",
                                            "{fixedJackpot}• Texto del jackpot"])

        self.ES[sentence] = "Juega hasta 50 líneas. ##sym_11.png##paga en cualquier posición."
        self.write()
        code, output = self.build("--lang", "esp")
        self.assertEqual(code, 0, output)
        self.assertNotIn("line_count_mismatch", output)
        # Sin saltos: la frase completa en la primera línea y la continuación eliminada.
        self.assertEqual(self.lines(), ["• Juega hasta 50 líneas. ##sym_11.png##paga en cualquier posición.", "",
                                        "{fixedJackpot}• Texto del jackpot", "    6 ##sym_12.png##."])
        report = json.loads((self.out / "xlsx-build-report.json").read_text(encoding="utf-8"))
        esp = report["languages"]["esp"]
        self.assertNotIn("T0002", esp["missing"])
        self.assertNotIn("T0003", esp["missing"])
        self.assertEqual(esp["merged_lines"], ["rules.pages[0].items.lines[2]"])
        self.assertEqual(esp["marker_issues"], [])
        eng = json.loads(self.template.read_text(encoding="utf-8"))
        esp_data = json.loads((self.out / "esp" / "info.json").read_text(encoding="utf-8"))
        del eng["rules"]["pages"][0]["items"]["lines"][2]
        self.assertEqual(eng["rules"]["pages"][1], esp_data["rules"]["pages"][1])
        self.assertEqual(eng["rules"]["settings"], esp_data["rules"]["settings"])

        # El CSV del idioma conserva la unión: rebuild --texts reproduce el mismo archivo.
        merged = {row[0]: row[2] for row in read_csv(self.out / "sheets" / "esp.csv")[1:]}
        self.assertEqual(merged["T0002"], "T0003")
        built = (self.out / "esp" / "info.json").read_bytes()
        again = Path(self.tmp.name) / "again"
        code, output = run("rebuild", str(self.out / "manifest.json"), "--texts", str(self.out / "sheets" / "esp.csv"),
                           "--lang", "esp", "--out", str(again))
        self.assertEqual(code, 0, output)
        self.assertEqual((again / "esp" / "info.json").read_bytes(), built)

    def test_merged_sentence_that_loses_a_marker_stays_in_english(self):
        sentence = "Play up to 50 lines. ##sym_11.png##pays anywhere."
        self.ENG = ["PAYTABLE", sentence]  # la fila 1 es un rótulo sin fórmula, como en la hoja real
        self.ES = {sentence: "Juega hasta 50 líneas, paga en cualquier posición."}
        self.write()
        code, output = self.build("--lang", "esp")
        self.assertEqual(code, 0, output)
        self.assertEqual(self.lines()[:3], ["• Play up to 50 lines. ", "", "   ##sym_11.png##pays anywhere."])
        esp = json.loads((self.out / "xlsx-build-report.json").read_text(encoding="utf-8"))["languages"]["esp"]
        self.assertEqual([(i["key"], i["kind"]) for i in esp["marker_issues"]], [("T0002", "token_mismatch")])
        self.assertEqual(esp["merged_lines"], [])
        self.assertIn("T0003", esp["missing"])

    def test_same_english_twice_needs_the_same_translation(self):
        rows = ["GAME RULES", "OF A KIND", "OF A KIND"]
        eng = [["", "", t] for t in rows]
        for second, expected in (("DE UN TIPO", 0), ("DEL MISMO TIPO", 1)):
            values = ["REGLAS", "DE UN TIPO", second]
            esp = [["", "", F(f"=ENG!C{r}", v)] for r, v in enumerate(values, start=1)]
            write_xlsx(self.xlsx, {"ENG": eng, "ESP": esp})
            code, output = self.build("--lang", "esp")
            self.assertEqual(code, expected, output)
        self.assertIn("ambiguous", output)
        self.assertIn("traducciones distintas", output)


@unittest.skipUnless(HAS_OPENPYXL, "openpyxl no está instalado")
class GenerateCommandTest(TextModeCase):
    """`generate` = extract + xlsx build en un solo paso, con la misma salida que por separado."""

    def generate(self, *extra):
        return run("generate", str(self.template), "--out", str(self.out), "--sheets-config", str(TEXT_MODE), *extra)

    def test_generate_matches_extract_plus_build(self):
        self.write()
        self.assertEqual(self.build()[0], 0)
        separate = {p.parent.name: p.read_bytes() for p in self.out.glob("*/info.json")}
        for path in self.out.glob("*/info.json"):
            path.unlink()
        code, output = self.generate("--xlsx", str(self.xlsx))
        self.assertEqual(code, 0, output)
        self.assertIn("Claves respecto al manifest anterior: unchanged=9", output)
        self.assertEqual({p.parent.name: p.read_bytes() for p in self.out.glob("*/info.json")}, separate)
        self.assertEqual(sorted(separate), ["deu", "eng", "esp"])

    def test_generate_requires_the_xlsx_before_extracting(self):
        manifest = self.out / "manifest.json"
        before = manifest.read_bytes()
        code, output = self.generate()
        self.assertEqual(code, 2)
        self.assertIn("--xlsx", output)
        self.assertEqual(manifest.read_bytes(), before)

    def test_version(self):
        with redirect_stdout(io.StringIO()) as out, self.assertRaises(SystemExit) as exit_:
            main(["--version"])
        self.assertEqual(exit_.exception.code, 0)
        self.assertRegex(out.getvalue(), r"^translateml \d+\.\d+\.\d+")


@unittest.skipUnless(HAS_OPENPYXL and os.environ.get("TRANSLATEML_REAL_XLSX"), "define TRANSLATEML_REAL_XLSX con la ruta del archivo real")
class RealXlsxTest(unittest.TestCase):
    """Integración con el Multidioma.xlsx real (g490) usando config/sheets.json; no modifica el archivo."""

    def test_builds_g490_from_real_file(self):
        real = Path(os.environ["TRANSLATEML_REAL_XLSX"]).expanduser()
        before = real.read_bytes()
        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp)
            g490 = next(f for f in REAL_TEMPLATES if f.parent.name == "g490")
            self.assertEqual(run("extract", str(g490), "--out", str(out))[0], 0)
            with redirect_stdout(io.StringIO()) as log:
                code = main(["xlsx", "build", str(out / "manifest.json"), "--xlsx", str(real)])
            self.assertEqual(code, 0, log.getvalue())
            report = json.loads((out / "xlsx-build-report.json").read_text(encoding="utf-8"))
            self.assertNotEqual(report["status"], "error")
            self.assertEqual((out / "eng" / "info.json").read_bytes(), g490.read_bytes())
            for code_ in ("esp", "deu"):
                self.assertGreater(report["languages"][code_]["applied"], 0)
                json.loads((out / code_ / "info.json").read_text(encoding="utf-8"))
        self.assertEqual(real.read_bytes(), before)


if __name__ == "__main__":
    unittest.main()
