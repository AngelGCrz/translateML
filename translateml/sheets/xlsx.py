"""Lectura de solo lectura de un .xlsx descargado de Google Sheets (Multidioma.xlsx).

openpyxl no calcula fórmulas: se leen los valores que Google Sheets guardó al exportar
(data_only=True) y, por separado, el texto de las fórmulas, que solo sirve para saber a
qué fila de ENG apunta cada traducción. Una fórmula sin valor guardado se registra para
informarla; nunca se usa la fórmula como traducción. El archivo se lee entero en memoria
y no se escribe.
"""
from __future__ import annotations

import datetime
import difflib
import hashlib
import zipfile
from pathlib import Path as FilePath
from typing import Dict, List, Optional, Set, Tuple

from .errors import SheetsReadError, SheetsStructureError

_INSTALL = ("Falta openpyxl para leer archivos .xlsx. Instálalo con:\n"
            "  python3 -m venv .venv && .venv/bin/python -m pip install -r requirements.txt")


class XlsxWorkbook:
    def __init__(self, file):
        self.file = FilePath(file)
        if not self.file.is_file():
            raise SheetsReadError(f"No existe el archivo {self.file}.")
        if self.file.suffix.lower() not in (".xlsx", ".xlsm"):
            raise SheetsReadError(f"{self.file.name} no es un .xlsx. Descárgalo desde Google Sheets con "
                                  "Archivo > Descargar > Microsoft Excel (.xlsx).")
        try:
            import openpyxl
            from openpyxl.utils.exceptions import InvalidFileException
        except ImportError as exc:
            raise SheetsReadError(_INSTALL) from exc
        raw = self.file.read_bytes()
        self.sha256 = hashlib.sha256(raw).hexdigest()
        try:
            computed = openpyxl.load_workbook(self.file, read_only=True, data_only=True)
            written = openpyxl.load_workbook(self.file, read_only=True, data_only=False)
        except (zipfile.BadZipFile, InvalidFileException, KeyError, OSError, ValueError) as exc:
            raise SheetsReadError(f"No se pudo leer {self.file.name} como XLSX: {exc}") from exc
        try:
            self._titles = list(computed.sheetnames)
            self._values: Dict[str, List[List[str]]] = {}
            self._formulas: Dict[str, List[List[str]]] = {}
            self._uncached: Dict[str, Set[Tuple[int, int]]] = {}
            for title in self._titles:
                self._load(title, computed[title], written[title])
        finally:
            computed.close()
            written.close()

    def _load(self, title, computed, written) -> None:
        values = [[_text(v) for v in row] for row in computed.iter_rows(values_only=True)]
        formulas, uncached = [], set()
        for r, row in enumerate(written.iter_rows(values_only=True), start=1):
            cells = []
            for c, raw in enumerate(row):
                formula = _formula(raw)
                if formula is None:
                    cells.append(_text(raw))
                    continue
                cells.append(formula)
                if r > len(values) or c >= len(values[r - 1]) or values[r - 1][c] == "":
                    uncached.add((r, c))
            formulas.append(cells)
        self._values[title], self._formulas[title], self._uncached[title] = values, formulas, uncached

    def tabs(self) -> List[str]:
        return list(self._titles)

    def read(self, tab: str, formulas: bool = False) -> List[List[str]]:
        if tab not in self._values:
            raise SheetsStructureError(missing_tab_message(self.file.name, tab, self._titles))
        return [list(row) for row in (self._formulas if formulas else self._values)[tab]]

    def missing_values(self, tab: str) -> Set[Tuple[int, int]]:
        """Celdas con fórmula y sin valor guardado: (fila 1-based, columna 0-based)."""
        return set(self._uncached.get(tab, set()))


def suggest_tab(name: str, tabs: List[str]) -> Optional[str]:
    """Pestaña existente que probablemente sea `name` (otras mayúsculas, espacios o una errata)."""
    by_lower = {t.strip().lower(): t for t in tabs}
    if name.strip().lower() in by_lower:
        return by_lower[name.strip().lower()]
    close = difflib.get_close_matches(name, tabs, n=1, cutoff=0.75)
    return close[0] if close else None


def missing_tab_message(file_name: str, tab: str, tabs: List[str]) -> str:
    hint = suggest_tab(tab, tabs)
    hint = f" ¿Es {hint!r}? Si es así, corrige 'tab' en config/languages.json." if hint else ""
    return f"{file_name} no tiene la pestaña {tab!r} (pestañas: {', '.join(tabs)}).{hint}"


def _formula(raw):
    text = getattr(raw, "text", raw)  # ArrayFormula de openpyxl expone el texto en .text
    if isinstance(text, str) and text.startswith("="):
        return text
    return None


def _text(value) -> str:
    """Valor de celda como texto. None (celda vacía) -> "", distinto de una cadena de espacios."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, (datetime.datetime, datetime.date, datetime.time)):
        return value.isoformat()
    return str(value)
